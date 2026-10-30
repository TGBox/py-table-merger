# type: ignore
"""Tests für den Lemniscus-Export (Importformat laut Kundenvorlage)."""
import datetime
import os
import sys
import unittest

import openpyxl
import pandas as pd

current_dir = os.path.dirname(os.path.abspath(__file__))
root_dir = os.path.abspath(os.path.join(current_dir, ".."))
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)
if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

from src.lemniscus import lemniscus_export as lx
from test_patient_record_export import PatientRecordTestBase, write_csv

TEMPLATE = os.path.join(current_dir, "fixtures", "lemniscus", "akten-vorlage.xlsx")


def _column_widths(path) -> dict:
    """(Blattname, Spaltennummer) -> Breite aus den <col>-Elementen der xlsx-Datei."""
    import re
    import zipfile
    wb = openpyxl.load_workbook(path)
    result = {}
    with zipfile.ZipFile(path) as z:
        for i, name in enumerate(wb.sheetnames, start=1):
            xml = z.read(f"xl/worksheets/sheet{i}.xml").decode("utf-8")
            for lo, hi, width in re.findall(r'<col min="(\d+)" max="(\d+)" width="([\d.]+)"', xml):
                for col in range(int(lo), int(hi) + 1):
                    result[(name, col)] = float(width)
    return result


class TestHelpers(unittest.TestCase):
    def test_normalize_phone(self):
        self.assertEqual(lx.normalize_phone("0451 27798"), "+4945127798")
        self.assertEqual(lx.normalize_phone("0175 233 0733"), "+491752330733")
        self.assertEqual(lx.normalize_phone("0049 172/123-45"), "+4917212345")
        self.assertEqual(lx.normalize_phone("+49 (0)172 12345"), "+4917212345")
        self.assertEqual(lx.normalize_phone("880332"), "880332")  # ohne Vorwahl: nicht raten
        self.assertEqual(lx.normalize_phone("nur abends"), "nur abends")
        self.assertIsNone(lx.normalize_phone(None))

    def test_street_with_number(self):
        self.assertEqual(lx.street_with_number("Kohlenhof", "7"), "Kohlenhof 7")
        self.assertEqual(lx.street_with_number("Kohlenhof 7", "7"), "Kohlenhof 7")
        self.assertEqual(lx.street_with_number("Kohlenhof", None), "Kohlenhof")

    def test_diagnosis_text_falls_back_to_icd(self):
        self.assertEqual(lx.diagnosis_text({"diagnose": "Gangstörung", "icd10_1": "R26.2"}), "Gangstörung")
        self.assertEqual(lx.diagnosis_text({"diagnose": None, "icd10_1": "M54.2", "icd10_2": "R52.2"}),
                         "M54.2, R52.2")
        self.assertIsNone(lx.diagnosis_text({}))

    def test_parse_recipient_json_person(self):
        r = lx.parse_recipient('[{"anrede":"Herr","name":"Ehlers Carl-Christian","plz":"24226",'
                               '"ort":"Heikendorf","strassepostfach":"Schlosskoppelweg 33"}]')
        self.assertEqual(r, {"Rechnung Anrede": "Herr", "Rechnung Nachname": "Ehlers",
                             "Rechnung Vorname": "Carl-Christian", "Rechnung Straße": "Schlosskoppelweg 33",
                             "Rechnung PLZ": "24226", "Rechnung Stadt": "Heikendorf", "Rechnung Land": "DE"})

    def test_parse_recipient_json_institution(self):
        r = lx.parse_recipient('[{"anrede":"Institution","vorname":"ZDPol Heilfürsorge","plz":"14471",'
                               '"ort":"Potsdam","strassepostfach":"An der Pirschheide 11"}]')
        self.assertEqual(r["Rechnung Firma"], "ZDPol Heilfürsorge")
        self.assertNotIn("Rechnung Nachname", r)

    def test_parse_recipient_text(self):
        r = lx.parse_recipient("Irene Gahn \nc/o Pflegeheim\nZiegenhorst 25\n23570 Lübeck")
        self.assertEqual(r["Rechnung Vorname"], "Irene")
        self.assertEqual(r["Rechnung Nachname"], "Gahn")
        self.assertEqual(r["Rechnung Straße"], "Ziegenhorst 25")
        self.assertEqual(r["Rechnung Zusatz"], "c/o Pflegeheim")
        self.assertEqual((r["Rechnung PLZ"], r["Rechnung Stadt"]), ("23570", "Lübeck"))

    def test_parse_recipient_unclear(self):
        r = lx.parse_recipient("Tochter zahlt")
        self.assertTrue(r.get("_unklar"))
        self.assertEqual(lx.parse_recipient(None), {})


class TestLemniscusExport(PatientRecordTestBase):
    def setUp(self):
        super().setUp()
        # Termin eines gelöschten Patienten (wird übersprungen, aber gezählt)
        path = os.path.join(self.folder, "events.csv")
        with open(path, encoding="utf-8") as f:
            content = f.read()
        content += ('"E4";"2026-02-01 09:00:00";"2026-02-01 09:20:00";"20";"KG";"R1";"P-GELOESCHT";'
                    '"1";NULL;NULL;NULL;NULL;NULL;NULL;NULL;NULL;NULL\n')
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        self.out = os.path.join(self.folder, "lemniscus.xlsx")

    def export(self):
        return lx.export_lemniscus(self.folder, self.out)

    def sheets(self):
        tables, _ = lx.pr.load_export_folder(self.folder)
        return lx.build_lemniscus(tables)[0]

    # --- Format laut Vorlage -------------------------------------------------
    def test_matches_customer_template(self):
        """Blätter, Reihenfolge, Kopfzeilen und Kopf-Formatierung exakt wie die Kundenvorlage."""
        self.export()
        template = openpyxl.load_workbook(TEMPLATE)
        out = openpyxl.load_workbook(self.out)
        self.assertEqual(out.sheetnames[:3], template.sheetnames)
        self.assertEqual(out.sheetnames[3], lx.SHEET_INSURANCE)  # Zusatzblatt am Ende
        for ws in template.worksheets:
            o = out[ws.title]
            self.assertEqual([c.value for c in o[1]], [c.value for c in ws[1]], ws.title)
            for col in range(1, ws.max_column + 1):
                t, e = ws.cell(1, col), o.cell(1, col)
                self.assertEqual((e.font.name, e.font.sz, e.font.b), (t.font.name, t.font.sz, t.font.b))
                self.assertEqual(e.fill.fgColor.rgb, t.fill.fgColor.rgb, f"{ws.title} {t.value}")
        # Spaltenbreiten direkt aus dem XML (Bereiche wie min=5 max=6 berücksichtigen)
        expected, actual = _column_widths(TEMPLATE), _column_widths(self.out)
        for (sheet, col), width in expected.items():
            self.assertAlmostEqual(actual.get((sheet, col), 0), width, delta=0.2,
                                   msg=f"{sheet} Spalte {col}")

    def test_cell_types(self):
        self.export()
        wb = openpyxl.load_workbook(self.out)
        akten, notizen, termine = wb["Akten"], wb["Notizen"], wb["Termine"]
        head = {c.value: c.column for c in akten[1]}
        row = 3  # Öztürk (Adam in Zeile 2 hat ein Geburtsdatum vor 1900 -> leer)
        self.assertIsInstance(akten.cell(row, head["Nummer"]).value, int)
        self.assertIsInstance(akten.cell(row, head["Geburtsdatum"]).value, datetime.datetime)
        self.assertEqual(akten.cell(row, head["Geburtsdatum"]).number_format, "dd.mm.yyyy")
        self.assertIs(akten.cell(row, head["Firmenkontakt"]).value, False)
        self.assertIs(akten.cell(row, head["Aktiv"]).value, True)
        self.assertIsInstance(akten.cell(row, head["PLZ"]).value, str)  # führende Null bleibt
        self.assertEqual(termine["C2"].number_format, "dd.mm.yyyy hh:mm")
        self.assertIsInstance(termine["G2"].value, bool)
        self.assertIsInstance(termine["D2"].value, int)
        self.assertEqual(notizen["C2"].number_format, "dd.mm.yyyy hh:mm")
        self.assertEqual(notizen["D2"].number_format, "@")

    # --- Inhalte -------------------------------------------------------------
    def test_files(self):
        a = self.sheets()["Akten"].set_index("Nachname")
        self.assertEqual(a.loc["Adam", "Nummer"], 1)
        self.assertEqual(a.loc["Öztürk", "Nummer"], 2)
        oez = a.loc["Öztürk"]
        self.assertEqual(oez["Geschlecht"], "w")
        self.assertEqual(oez["Straße"], "Hauptstr. 5a")
        self.assertEqual(oez["PLZ"], "01067")
        self.assertEqual(oez["Land"], "DE")
        self.assertEqual(oez["Telefon Privat"], "+49351123")
        self.assertEqual(oez["Akten-Typ"], "Patient")
        self.assertEqual(oez["Status"], "Aktiv")
        self.assertTrue(pd.isna(oez["Diagnose"]) or oez["Diagnose"] is None)  # Verordnung ohne Diagnose/ICD
        adam = a.loc["Adam"]
        self.assertEqual(adam["Geschlecht"], "m")
        self.assertEqual(adam["Rechnung Nachname"], "Sohn")
        self.assertEqual(adam["Rechnung Stadt"], "Lübeck")
        self.assertTrue(pd.isna(adam["Geburtsdatum"]))  # 1890 kann Excel nicht -> leer

    def test_events(self):
        tables, _ = lx.pr.load_export_folder(self.folder)
        numbers = lx.file_numbers(tables)
        events, skipped = lx.build_events(tables, numbers)
        self.assertEqual(skipped, 1)                 # gelöschter Patient
        self.assertEqual(len(events), 2)             # Flexzeit ohne Patient fällt weg
        self.assertEqual(events["Akte Nummer"].tolist(), [2, 2])
        first, second = events.iloc[0], events.iloc[1]
        self.assertEqual(first["Beginn"], datetime.datetime(2026, 1, 2, 8, 0))
        self.assertFalse(first["Wahrgenommen"])
        self.assertIn("Nicht erschienen", first["Anmerkungen"])
        self.assertEqual(second["Terminbezeichnung"], "KG")
        self.assertEqual(second["Dauer"], 20)
        self.assertEqual(second["Raum"], 1)
        self.assertTrue(second["Fertig bearbeitet"])  # abgerechnet
        self.assertIn("Therapeut: Uwe Therapeut", second["Anmerkungen"])
        self.assertIn("Verordnung 152, Termin 1/6", second["Anmerkungen"])

    def test_event_diagnosis_from_prescription(self):
        path = os.path.join(self.folder, "rezepte.csv")
        with open(path, encoding="utf-8") as f:
            header, *rows = f.read().splitlines()
        write_csv(self.folder, "rezepte", [
            {"id": "Z1", "patienten_id": "P1", "rezeptnr": "152", "datum": "2025-12-20",
             "diagnose": None, "icd10_1": "M54.2", "icd10_2": None},
        ])
        tables, _ = lx.pr.load_export_folder(self.folder)
        events, _ = lx.build_events(tables, lx.file_numbers(tables))
        self.assertEqual(events["Diagnose"].tolist(), ["M54.2", "M54.2"])
        files, _ = lx.build_files(tables, lx.file_numbers(tables))
        self.assertEqual(files.set_index("Nachname").loc["Öztürk", "Diagnose"], "M54.2")

    def test_notes(self):
        n = self.sheets()["Notizen"]
        doc = n[n["Titel"] == lx.NOTE_TITLE_EVENT_DOC].iloc[0]
        self.assertEqual(doc["Akte Nummer"], 2)
        self.assertEqual(doc["Zeitpunkt"], datetime.datetime(2026, 1, 5, 10, 20))  # = Terminbeginn
        self.assertEqual(doc["Text"], "Stichwort: Befund\n\nMobilisation HWS")
        finding = n[n["Titel"] == "Befund: Fraktur"].iloc[0]
        self.assertIn("Ansicht: Muskulatur vorne (weiblich)", finding["Text"])
        questionnaire = n[n["Titel"] == "Fragebogen Anamnese"].iloc[0]
        self.assertEqual(questionnaire["Akte Nummer"], 1)
        self.assertEqual(questionnaire["Text"], "Schmerzstärke: 4\nRauchen Sie?: Ja")
        history = n[n["Titel"] == "Telefon"].iloc[0]
        self.assertEqual((history["Akte Nummer"], history["Text"]), (1, "Rückruf"))
        report = n[n["Titel"] == "Therapiebericht Verordnung 152"].iloc[0]
        self.assertEqual(report["Akte Nummer"], 2)
        # Praxis-Notiz ohne Patientenbezug fehlt
        self.assertNotIn("Praxisnotiz", " ".join(n["Text"]))

    def test_insurance_sheet(self):
        v = self.sheets()[lx.SHEET_INSURANCE].set_index("Nachname")
        self.assertTrue(v.loc["Adam", "Privatversichert"])
        self.assertTrue(v.loc["Öztürk", "Zuzahlungsbefreit"])
        self.assertEqual(v.loc["Öztürk", "Akte Nummer"], 2)

    def test_numbers_link_the_right_patient_across_sheets(self):
        s = self.sheets()
        files = s["Akten"].set_index("Nummer")
        for sheet in ("Notizen", "Termine", lx.SHEET_INSURANCE):
            col = "Akte Nummer"
            self.assertTrue(set(s[sheet][col]) <= set(files.index), sheet)
        oez_nr = files.index[files["Nachname"] == "Öztürk"][0]
        self.assertEqual(set(s["Termine"]["Akte Nummer"]), {oez_nr})

    def test_no_database_ids_and_no_formulas(self):
        import zipfile
        path = os.path.join(self.folder, "events.csv")
        with open(path, encoding="utf-8") as f:
            content = f.read().replace('"Mobilisation HWS"', '"=> deutliche Besserung"')
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        self.export()
        tables, _ = lx.pr.load_export_folder(self.folder)
        ids = set()
        for df in tables.values():
            for col in df.columns:
                if col in ("id", "p_nr", "refid", "resource", "id_patient", "id_region") or col.endswith("_id"):
                    ids |= set(df[col].dropna().astype(str))
        wb = openpyxl.load_workbook(self.out)
        leaks = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows()
                 for c in row if c.value is not None and str(c.value) in ids]
        self.assertEqual(leaks, [])
        with zipfile.ZipFile(self.out) as z:
            for name in z.namelist():
                if name.startswith("xl/worksheets/"):
                    self.assertNotIn("<f>", z.read(name).decode("utf-8"))
        texts = [c.value for c in wb["Notizen"]["D"]]
        self.assertTrue(any(str(t).endswith("=> deutliche Besserung") for t in texts))

    def test_required_fields_fallbacks_and_check(self):
        """Pflichtfelder (blau in der Vorlage): Dauer aus Ende-Beginn, fehlende Bezeichnung -> "Termin",
        verbleibende Lücken werden gemeldet."""
        path = os.path.join(self.folder, "events.csv")
        with open(path, encoding="utf-8") as f:
            content = f.read()
        content = content.replace('"2026-01-05 10:40:00";"20";"KG"', '"2026-01-05 10:45:00";NULL;NULL')
        content = content.replace('"E2";"2026-01-02 08:00:00"', '"E2";"kaputt"')
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        result = self.export()
        tables, _ = lx.pr.load_export_folder(self.folder)
        events, _ = lx.build_events(tables, lx.file_numbers(tables))
        fixed = events[events["Dauer"] == 25].iloc[0]
        self.assertEqual(fixed["Terminbezeichnung"], lx.DEFAULT_EVENT_TITLE)
        self.assertEqual(result.missing_required, {"Termine": {"Beginn": 1}})
        self.assertTrue(any("Pflichtfeld" in h and "Beginn" in h for h in result.hints))

    def test_result_hints(self):
        result = self.export()
        self.assertEqual(result.sheet_counts["Termine"], 2)
        self.assertEqual(result.skipped_events, 1)
        self.assertTrue(any("übersprungen" in h for h in result.hints))


if __name__ == "__main__":
    unittest.main()
